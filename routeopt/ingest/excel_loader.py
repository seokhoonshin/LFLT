"""Read the planner's shipment workbook and map its columns to the 4.1 schema.

Header matching is case-insensitive and whitespace-trimmed; common aliases are accepted.
When required headers cannot be matched automatically the UI shows a column-mapping step
and calls :func:`apply_mapping` with the planner's choices.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

CANONICAL_COLUMNS: list[str] = [
    "shipment_id",
    "purchase_order",
    "delivery_date",
    "order_type",
    "customer_name",
    "phone",
    "address",
    "city",
    "state",
    "zip",
    "latitude",
    "longitude",
    "model_code",
    "pieces",
    "weight_lb",
    "volume_cuft",
    "service_minutes",
    "notes",
    "origin_dc",
    "pallet_required",
    "priority",
    "memo",
    "declared_value",
    "bill_to_code",
]

REQUIRED_COLUMNS: list[str] = [
    "shipment_id",
    "delivery_date",
    "order_type",
    "customer_name",
    "address",
    "city",
    "state",
    "zip",
    "volume_cuft",
]

COLUMN_ALIASES: dict[str, list[str]] = {
    "shipment_id": ["shipment id", "shipment", "shipment no", "shipment number", "id"],
    "purchase_order": ["po", "po number", "purchase order", "po no", "order number", "order no"],
    "delivery_date": [
        "delivery date",
        "date",
        "del date",
        "delivery dt",
        "scheduled date",
        "due date",
    ],
    "order_type": ["order type", "type", "channel", "customer type"],
    "customer_name": ["customer", "customer name", "name", "consignee", "consignee name"],
    "phone": ["phone number", "telephone", "tel", "customer phone", "phone no", "mobile"],
    "address": ["street", "street address", "address 1", "address1", "addr", "address line 1"],
    "city": ["town", "city name"],
    "state": ["st", "province", "state code"],
    "zip": ["zipcode", "zip code", "postal code", "postal", "zip5"],
    "latitude": ["lat", "y"],
    "longitude": ["lon", "lng", "long", "x"],
    "model_code": ["model", "model code", "sku", "model no", "model number", "item"],
    "pieces": ["pcs", "qty", "quantity", "piece count", "pieces count", "units"],
    "weight_lb": ["weight", "weight lb", "weight lbs", "wt", "lbs", "weight (lb)", "total weight"],
    "volume_cuft": [
        "volume",
        "cuft",
        "cu ft",
        "cubic feet",
        "vol",
        "volume (cuft)",
        "total volume",
    ],
    "service_minutes": [
        "service",
        "service min",
        "service time",
        "install minutes",
        "service mins",
        "minutes",
        "install time",
    ],
    "notes": ["note", "comments", "remarks", "instructions", "delivery notes"],
    "origin_dc": [
        "dc",
        "origin dc",
        "origin",
        "from dc",
        "ship from",
        "frst shpg loc cd",
        "first shpg loc cd",
        "warehouse",
    ],
    "pallet_required": [
        "pallet",
        "pallet required",
        "pallet yn",
        "needs pallet",
        "pallet flag",
        "hq",
        "product group",
        "prod group",
    ],
    "priority": ["prio", "priority group"],
    "memo": ["memo text", "memo txt", "internal memo"],
    "declared_value": ["declared val", "declared value", "dcld val", "dcld_val_dlr"],
    "bill_to_code": ["bill to", "bill to code", "billto", "bill to cd", "bill_to_cd", "bill_to"],
}


def normalize_header(header: Any) -> str:
    """Lower-case, trim, and collapse whitespace/punctuation so headers compare loosely."""
    text = "" if header is None else str(header)
    text = text.replace("_", " ").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


@dataclass
class RawTable:
    """Rows as read from the workbook, keyed by the original header text."""

    headers: list[str]
    rows: list[dict[str, Any]]
    row_numbers: list[int]
    sheet_name: str = ""
    source_name: str = ""


@dataclass
class ColumnMapping:
    mapping: dict[str, str]  # canonical -> original header
    missing_required: list[str] = field(default_factory=list)
    unmapped_headers: list[str] = field(default_factory=list)

    @property
    def complete(self) -> bool:
        return not self.missing_required


@dataclass
class RawRow:
    row_number: int
    values: dict[str, Any]  # canonical -> raw value (None if the column is absent)

    def get(self, key: str) -> Any:
        return self.values.get(key)


def read_table(source: str | Path | bytes | io.BytesIO, sheet_name: str | None = None) -> RawTable:
    """Read the first sheet (or ``sheet_name``) of an .xlsx, or a .csv, into a RawTable."""
    source_name = ""
    if isinstance(source, str | Path):
        path = Path(source)
        source_name = path.name
        if path.suffix.lower() == ".csv":
            return _read_csv(path.read_bytes(), source_name)
        data: bytes = path.read_bytes()
    elif isinstance(source, io.BytesIO):
        data = source.getvalue()
    else:
        data = source
    if data[:4] != b"PK\x03\x04":  # not a zip container -> try CSV
        return _read_csv(data, source_name)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
        header_row_idx: int | None = None
        headers: list[str] = []
        rows: list[dict[str, Any]] = []
        row_numbers: list[int] = []
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if header_row_idx is None:
                if row is None or all(c is None or str(c).strip() == "" for c in row):
                    continue
                header_row_idx = idx
                headers = [
                    str(c).strip() if c is not None else f"column_{i + 1}"
                    for i, c in enumerate(row)
                ]
                continue
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            values = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
            rows.append(values)
            row_numbers.append(idx)
        return RawTable(
            headers=headers,
            rows=rows,
            row_numbers=row_numbers,
            sheet_name=ws.title,
            source_name=source_name,
        )
    finally:
        wb.close()


def _read_csv(data: bytes, source_name: str) -> RawTable:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    headers: list[str] = []
    rows: list[dict[str, Any]] = []
    row_numbers: list[int] = []
    for idx, row in enumerate(reader, start=1):
        if not headers:
            if not row or all(not c.strip() for c in row):
                continue
            headers = [c.strip() or f"column_{i + 1}" for i, c in enumerate(row)]
            continue
        if not row or all(not c.strip() for c in row):
            continue
        values = {headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))}
        rows.append({k: (v if v != "" else None) for k, v in values.items()})
        row_numbers.append(idx)
    return RawTable(
        headers=headers,
        rows=rows,
        row_numbers=row_numbers,
        sheet_name="csv",
        source_name=source_name,
    )


def auto_map_columns(headers: list[str]) -> ColumnMapping:
    """Match workbook headers to canonical column names (exact first, then aliases)."""
    normalized = {normalize_header(h): h for h in headers}
    mapping: dict[str, str] = {}
    used: set[str] = set()
    # exact canonical matches first
    for canonical in CANONICAL_COLUMNS:
        key = normalize_header(canonical)
        if key in normalized and normalized[key] not in used:
            mapping[canonical] = normalized[key]
            used.add(normalized[key])
    # then aliases
    for canonical in CANONICAL_COLUMNS:
        if canonical in mapping:
            continue
        for alias in COLUMN_ALIASES.get(canonical, []):
            key = normalize_header(alias)
            if key in normalized and normalized[key] not in used:
                mapping[canonical] = normalized[key]
                used.add(normalized[key])
                break
    missing = [c for c in REQUIRED_COLUMNS if c not in mapping]
    unmapped = [h for h in headers if h not in used]
    return ColumnMapping(mapping=mapping, missing_required=missing, unmapped_headers=unmapped)


def apply_mapping(table: RawTable, mapping: dict[str, str]) -> list[RawRow]:
    """Project the raw rows onto canonical column names using ``mapping``."""
    out: list[RawRow] = []
    for values, row_number in zip(table.rows, table.row_numbers, strict=True):
        projected = {canonical: values.get(original) for canonical, original in mapping.items()}
        for canonical in CANONICAL_COLUMNS:
            projected.setdefault(canonical, None)
        out.append(RawRow(row_number=row_number, values=projected))
    return out


def load_rows(
    source: str | Path | bytes | io.BytesIO, mapping: dict[str, str] | None = None
) -> tuple[RawTable, ColumnMapping, list[RawRow]]:
    """Convenience: read + auto-map (+ optional manual mapping override) + project."""
    table = read_table(source)
    column_mapping = auto_map_columns(table.headers)
    if mapping:
        merged = dict(column_mapping.mapping)
        merged.update({k: v for k, v in mapping.items() if v})
        column_mapping = ColumnMapping(
            mapping=merged,
            missing_required=[c for c in REQUIRED_COLUMNS if c not in merged],
            unmapped_headers=[h for h in table.headers if h not in merged.values()],
        )
    rows = apply_mapping(table, column_mapping.mapping)
    return table, column_mapping, rows


def excel_serial_to_date(value: float) -> dt.date:
    """Convert an Excel serial date number to a date (1900 date system)."""
    return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(value))).date()
