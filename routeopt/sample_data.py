"""Synthetic shipment workbook generator (brief section 16).

Produces an Excel file in the 4.1 schema: coordinates scattered within 40 miles of the hub,
~70/20/10 OBS/BUILDER/LTL, ~80/12/8 LMD/RTN/INC, realistic LG product volumes and service
times, shared addresses and POs, dates across D/D+1/D+2, a few rows with blank lat/lon that
use real public addresses (so the live Census geocoder can match them), and one builder
stop near 600 cu ft to exercise multi-load.
"""

from __future__ import annotations

import datetime as dt
import math
import random
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from routeopt.domain import HUB_LAT, HUB_LON
from routeopt.ingest.excel_loader import CANONICAL_COLUMNS

# (city, zip, lat, lon) - Inland Empire, east LA County and north Orange County places
# within ~40 miles of the Rialto hub (city/ZIP centroids; the generator jitters each by up
# to 3 miles).
CITIES: list[tuple[str, str, float, float]] = [
    ("Rialto", "92376", 34.1064, -117.3703),
    ("Rialto", "92377", 34.1480, -117.4050),
    ("Fontana", "92335", 34.0922, -117.4350),
    ("Fontana", "92336", 34.1340, -117.4620),
    ("Bloomington", "92316", 34.0703, -117.3959),
    ("San Bernardino", "92401", 34.1083, -117.2898),
    ("San Bernardino", "92404", 34.1470, -117.2560),
    ("San Bernardino", "92407", 34.1950, -117.3350),
    ("Colton", "92324", 34.0739, -117.3136),
    ("Grand Terrace", "92313", 34.0339, -117.3136),
    ("Loma Linda", "92354", 34.0483, -117.2612),
    ("Redlands", "92373", 34.0556, -117.1825),
    ("Highland", "92346", 34.1283, -117.2086),
    ("Yucaipa", "92399", 34.0336, -117.0431),
    ("Rancho Cucamonga", "91730", 34.1064, -117.5931),
    ("Rancho Cucamonga", "91739", 34.1250, -117.5150),
    ("Ontario", "91764", 34.0633, -117.6509),
    ("Ontario", "91761", 34.0330, -117.6150),
    ("Upland", "91786", 34.0975, -117.6484),
    ("Montclair", "91763", 34.0775, -117.6897),
    ("Chino", "91710", 34.0122, -117.6889),
    ("Chino Hills", "91709", 33.9898, -117.7326),
    ("Pomona", "91766", 34.0551, -117.7523),
    ("Claremont", "91711", 34.0967, -117.7198),
    ("Riverside", "92501", 33.9806, -117.3755),
    ("Riverside", "92503", 33.9380, -117.4400),
    ("Riverside", "92507", 33.9800, -117.3300),
    ("Jurupa Valley", "92509", 34.0000, -117.4600),
    ("Eastvale", "92880", 33.9638, -117.5640),
    ("Norco", "92860", 33.9311, -117.5487),
    ("Corona", "92882", 33.8753, -117.5664),
    ("Moreno Valley", "92553", 33.9425, -117.2297),
    ("Perris", "92571", 33.7825, -117.2286),
    ("Beaumont", "92223", 33.9295, -116.9770),
    ("Banning", "92220", 33.9255, -116.8767),
    ("Hesperia", "92345", 34.4264, -117.3009),
    ("Victorville", "92392", 34.5121, -117.3236),
    ("Apple Valley", "92307", 34.5008, -117.1859),
    ("Lake Arrowhead", "92352", 34.2486, -117.1896),
    ("Diamond Bar", "91765", 34.0286, -117.8103),
    ("West Covina", "91790", 34.0686, -117.9390),
    ("Glendora", "91741", 34.1361, -117.8653),
    ("San Dimas", "91773", 34.1067, -117.8067),
    ("Brea", "92821", 33.9167, -117.9000),
    ("Yorba Linda", "92886", 33.8886, -117.8131),
    ("Lake Elsinore", "92530", 33.6681, -117.3273),
]

STREET_NAMES = [
    "Sierra",
    "Oak",
    "Maple",
    "Magnolia",
    "Dogwood",
    "Hickory",
    "Willow",
    "Cedar",
    "Lake",
    "Ridge",
    "Valley",
    "Meadow",
    "Forest",
    "River",
    "Creek",
    "Church",
    "Main",
    "Mill",
    "Jackson",
    "Jefferson",
    "Hillcrest",
    "Sunset",
    "Arrowhead",
    "Highland",
    "Foothill",
    "Baseline",
    "Etiwanda",
    "Euclid",
    "Fairview",
    "Mission",
]
STREET_SUFFIXES = ["Rd", "Dr", "Ln", "Ct", "Way", "Blvd", "Ave", "St", "Trl", "Cir", "Pkwy"]
FIRST_NAMES = [
    "James",
    "Mary",
    "Robert",
    "Patricia",
    "John",
    "Jennifer",
    "Michael",
    "Linda",
    "David",
    "Elizabeth",
    "William",
    "Barbara",
    "Richard",
    "Susan",
    "Joseph",
    "Jessica",
    "Thomas",
    "Sarah",
    "Daniel",
    "Karen",
    "Marcus",
    "Aaliyah",
    "Andre",
    "Keisha",
    "Luis",
    "Maria",
    "Wei",
    "Priya",
    "Chen",
    "Fatima",
    "Kenji",
    "Olga",
]
LAST_NAMES = [
    "Smith",
    "Johnson",
    "Williams",
    "Brown",
    "Jones",
    "Garcia",
    "Miller",
    "Davis",
    "Rodriguez",
    "Martinez",
    "Wilson",
    "Anderson",
    "Taylor",
    "Thomas",
    "Moore",
    "Jackson",
    "Lee",
    "Harris",
    "Clark",
    "Lewis",
    "Walker",
    "Young",
    "Allen",
    "King",
    "Wright",
    "Scott",
    "Green",
    "Baker",
    "Adams",
    "Nelson",
    "Carter",
    "Mitchell",
    "Patel",
    "Nguyen",
]
BUILDERS = [
    "Pulte Homes",
    "DR Horton",
    "KB Home",
    "Lennar Inland Empire",
    "Tri Pointe Homes",
    "Shea Homes",
]
LTL_CONSIGNEES = [
    "Inland Appliance Outlet",
    "Metro Builders Supply",
    "Golden State Furniture",
    "Riverside Rental Center",
    "San Bernardino County Housing Authority",
]
NOTES = [
    "",
    "",
    "",
    "",
    "Call 30 min before arrival",
    "Gate code 4321",
    "Haul away old unit",
    "Stairs - 2nd floor",
    "Dog in yard",
    "Deliver to side door",
    "Customer works from home",
    "Narrow driveway - park on street",
    "Water line shutoff in basement",
]

# Real, public addresses used for the blank lat/lon rows so the live Census geocoder matches them.
REAL_ADDRESSES: list[tuple[str, str, str, str]] = [
    ("385 N Arrowhead Ave", "San Bernardino", "CA", "92415"),  # SB County Government Center
    ("4080 Lemon St", "Riverside", "CA", "92501"),  # Riverside County Administrative Center
    ("8353 Sierra Ave", "Fontana", "CA", "92335"),  # Fontana City Hall
    ("303 E B St", "Ontario", "CA", "91764"),  # Ontario City Hall
    ("10500 Civic Center Dr", "Rancho Cucamonga", "CA", "91730"),  # Rancho Cucamonga City Hall
    ("5500 University Pkwy", "San Bernardino", "CA", "92407"),  # Cal State San Bernardino
]


@dataclass(frozen=True)
class Product:
    category: str
    model: str
    vol_range: tuple[float, float]
    weight_range: tuple[float, float]
    service_range: tuple[int, int]


PRODUCTS: list[Product] = [
    Product("TV", "OLED65C4PUA", (5, 15), (40, 80), (60, 90)),
    Product("TV", "OLED77C4PUA", (8, 15), (60, 90), (60, 90)),
    Product("TV", "65UT8000AUA", (5, 12), (40, 70), (60, 90)),
    Product("TV", "75QNED80TUA", (8, 15), (60, 90), (60, 90)),
    Product("TV", "86UT7570PUA", (10, 15), (70, 90), (70, 90)),
    Product("Refrigerator", "LRFXC2606S", (60, 90), (250, 350), (30, 60)),
    Product("Refrigerator", "LRMVS3006S", (65, 90), (280, 350), (30, 60)),
    Product("Refrigerator", "LRYKC2606S", (60, 85), (250, 330), (30, 60)),
    Product("Refrigerator", "LRSXS2706V", (60, 80), (250, 320), (30, 60)),
    Product("Range", "LSGL6337F", (30, 50), (150, 220), (30, 45)),
    Product("Range", "LREL6325F", (30, 45), (150, 210), (30, 45)),
    Product("Range", "LSEL6337F", (30, 50), (160, 220), (30, 45)),
    Product("Dishwasher", "LDTH7972S", (12, 20), (80, 110), (40, 60)),
    Product("Washer", "WM4000HBA", (25, 35), (180, 230), (30, 45)),
    Product("Dryer", "DLEX4000B", (25, 35), (120, 160), (30, 45)),
]
CATEGORY_WEIGHTS = {
    "TV": 0.35,
    "Refrigerator": 0.25,
    "Range": 0.2,
    "Dishwasher": 0.08,
    "Washer": 0.06,
    "Dryer": 0.06,
}


def _offset_point(
    rng: random.Random, lat: float, lon: float, max_miles: float
) -> tuple[float, float]:
    r = max_miles * math.sqrt(rng.random())
    theta = rng.random() * 2 * math.pi
    dlat = (r * math.cos(theta)) / 69.0
    dlon = (r * math.sin(theta)) / (69.0 * math.cos(math.radians(lat)))
    return round(lat + dlat, 6), round(lon + dlon, 6)


def _miles(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    r = 3958.7613
    p1, p2 = math.radians(lat1), math.radians(lat2)
    a = (
        math.sin((p2 - p1) / 2) ** 2
        + math.cos(p1) * math.cos(p2) * math.sin(math.radians(lon2 - lon1) / 2) ** 2
    )
    return 2 * r * math.asin(math.sqrt(a))


def _pick_product(rng: random.Random) -> Product:
    cat = rng.choices(list(CATEGORY_WEIGHTS), weights=list(CATEGORY_WEIGHTS.values()))[0]
    return rng.choice([p for p in PRODUCTS if p.category == cat])


@dataclass
class _Address:
    street: str
    city: str
    state: str
    zip: str
    lat: float | None
    lon: float | None
    order_type: str
    customer: str
    phone: str
    day_offset: int


def _make_address(rng: random.Random, order_type: str, day_offset: int) -> _Address:
    while True:
        city, zip_code, clat, clon = rng.choice(CITIES)
        lat, lon = _offset_point(rng, clat, clon, 3.0)
        if _miles(HUB_LAT, HUB_LON, lat, lon) <= 40:
            break
    street = f"{rng.randint(100, 9999)} {rng.choice(STREET_NAMES)} {rng.choice(STREET_SUFFIXES)}"
    if order_type == "OBS":
        customer = f"{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)}"
    elif order_type == "BUILDER":
        customer = f"{rng.choice(BUILDERS)} - Lot {rng.randint(1, 120)}"
    else:
        customer = rng.choice(LTL_CONSIGNEES)
    phone = f"(909) 555-{rng.randint(100, 199):03d}{rng.randint(0, 9)}"
    return _Address(street, city, "CA", zip_code, lat, lon, order_type, customer, phone, day_offset)


def generate_rows(
    n_shipments: int, plan_date: dt.date, seed: int = 42, blank_coords: bool = True
) -> list[dict[str, object]]:
    """Generate ``n_shipments`` rows in the 4.1 schema."""
    rng = random.Random(seed)
    rows: list[dict[str, object]] = []
    n_stops = max(3, int(n_shipments * 0.6))
    addresses: list[_Address] = []
    for _ in range(n_stops):
        order_type = rng.choices(["OBS", "BUILDER", "LTL"], weights=[70, 20, 10])[0]
        day_offset = rng.choices([0, 1, 2], weights=[70, 20, 10])[0]
        addresses.append(_make_address(rng, order_type, day_offset))

    po_counter = 1000
    ship_counter = 1

    def next_po() -> str:
        nonlocal po_counter
        po_counter += 1
        return f"PO{po_counter}"

    def next_id(suffix: str) -> str:
        nonlocal ship_counter
        sid = f"SH{ship_counter:05d}_{suffix}"
        ship_counter += 1
        return sid

    def base_row(addr: _Address, suffix: str, po: str, day_offset: int) -> dict[str, object]:
        return {
            "shipment_id": next_id(suffix),
            "purchase_order": po,
            "delivery_date": plan_date + dt.timedelta(days=day_offset),
            "order_type": addr.order_type,
            "customer_name": addr.customer,
            "phone": addr.phone,
            "address": addr.street,
            "city": addr.city,
            "state": addr.state,
            "zip": addr.zip,
            "latitude": addr.lat,
            "longitude": addr.lon,
            "model_code": None,
            "pieces": 0,
            "weight_lb": 0.0,
            "volume_cuft": 0.0,
            "service_minutes": 0,
            "notes": rng.choice(NOTES) or None,
        }

    def lmd_row(
        addr: _Address,
        po: str,
        day_offset: int,
        product: Product | None = None,
        pieces: int | None = None,
    ) -> dict[str, object]:
        p = product or _pick_product(rng)
        pcs = pieces or (rng.choices([1, 2], weights=[85, 15])[0] if p.category != "TV" else 1)
        row = base_row(addr, "LMD", po, day_offset)
        row.update(
            model_code=p.model,
            pieces=pcs,
            weight_lb=round(rng.uniform(*p.weight_range) * pcs, 1),
            volume_cuft=round(rng.uniform(*p.vol_range) * pcs, 1),
            service_minutes=int(rng.randint(*p.service_range) * (1 if pcs == 1 else 1.5)),
        )
        return row

    def rtn_row(addr: _Address, po: str, day_offset: int) -> dict[str, object]:
        p = _pick_product(rng)
        row = base_row(addr, "RTN", po, day_offset)
        row.update(
            model_code=p.model,
            pieces=1,
            weight_lb=round(rng.uniform(*p.weight_range), 1),
            volume_cuft=round(rng.uniform(*p.vol_range), 1),
            service_minutes=20,
        )
        return row

    def inc_row(addr: _Address, po: str, day_offset: int) -> dict[str, object]:
        row = base_row(addr, "INC", po, day_offset)
        row.update(
            model_code=rng.choice(PRODUCTS).model,
            pieces=0,
            weight_lb=0.0,
            volume_cuft=0.0,
            service_minutes=30,
        )
        return row

    # 1) one builder stop near 600 cu ft (multi-load exerciser)
    big = _make_address(rng, "BUILDER", 0)
    big.customer = "Pulte Homes - Building 7 (8 units)"
    po = next_po()
    fridge = next(p for p in PRODUCTS if p.model == "LRFXC2606S")
    rng_product = next(p for p in PRODUCTS if p.model == "LSGL6337F")
    for _ in range(6):
        r = lmd_row(big, po, 0, fridge, 1)
        r["volume_cuft"] = 80.0
        r["weight_lb"] = 300.0
        r["service_minutes"] = 20
        rows.append(r)
    for _ in range(3):
        r = lmd_row(big, po, 0, rng_product, 1)
        r["volume_cuft"] = 36.0
        r["weight_lb"] = 180.0
        r["service_minutes"] = 15
        rows.append(r)  # 6*80 + 3*36 = 588 cu ft

    # 2) every address gets at least one shipment
    for addr in addresses:
        if len(rows) >= n_shipments:
            break
        po = next_po()
        rows.append(lmd_row(addr, po, addr.day_offset))

    # 3) extra shipments on shared addresses / POs until we reach n
    while len(rows) < n_shipments:
        addr = rng.choice(addresses)
        kind = rng.choices(
            ["LMD_same_po", "LMD_new_po", "RTN", "INC", "LMD_next_day"], weights=[20, 15, 35, 22, 8]
        )[0]
        existing = [r for r in rows if r["address"] == addr.street and r["zip"] == addr.zip]
        po = str(existing[0]["purchase_order"]) if existing and kind == "LMD_same_po" else next_po()
        if kind == "RTN":
            rows.append(rtn_row(addr, po, addr.day_offset))
        elif kind == "INC":
            rows.append(inc_row(addr, po, addr.day_offset))
        elif kind == "LMD_next_day":
            rows.append(lmd_row(addr, po, min(2, addr.day_offset + 1)))
        else:
            rows.append(lmd_row(addr, po, addr.day_offset))

    # 4) a few rows with blank lat/lon on real addresses (geocoder exerciser)
    n_blank = min(4, max(1, n_shipments // 30)) if blank_coords else 0
    candidates = [r for r in rows if r["address"] != big.street]
    for i in range(min(n_blank, len(candidates))):
        street, city, state, zip_code = REAL_ADDRESSES[i % len(REAL_ADDRESSES)]
        r = rng.choice(candidates)
        candidates = [c for c in candidates if c["customer_name"] != r["customer_name"]]
        r["address"], r["city"], r["state"], r["zip"] = street, city, state, zip_code
        r["latitude"], r["longitude"] = None, None
        # keep other shipments at the old address consistent: give them the new address too
        for other in rows:
            if (
                other is not r
                and other["customer_name"] == r["customer_name"]
                and other["phone"] == r["phone"]
            ):
                other["address"], other["city"], other["state"], other["zip"] = (
                    street,
                    city,
                    state,
                    zip_code,
                )
                other["latitude"], other["longitude"] = None, None

    return rows[:n_shipments] if n_shipments >= 9 else rows


def write_workbook(rows: list[dict[str, object]], path: Path) -> Path:
    """Write rows to an .xlsx with the 4.1 headers; zip kept as text, dates as dates."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipments"
    ws.append(CANONICAL_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(c) for c in CANONICAL_COLUMNS])
    date_col = CANONICAL_COLUMNS.index("delivery_date") + 1
    zip_col = CANONICAL_COLUMNS.index("zip") + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=date_col).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=zip_col).number_format = "@"
    widths = {
        "shipment_id": 16,
        "purchase_order": 12,
        "delivery_date": 13,
        "customer_name": 30,
        "address": 26,
        "city": 14,
        "notes": 30,
        "model_code": 14,
    }
    for i, c in enumerate(CANONICAL_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 11)
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def generate_sample(
    n_shipments: int, plan_date: dt.date, seed: int, out: Path, blank_coords: bool = True
) -> Path:
    return write_workbook(generate_rows(n_shipments, plan_date, seed, blank_coords), out)
