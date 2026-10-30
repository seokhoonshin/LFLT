"""Appointment tracking: the trucks the planner picks on Results, one row per shipment.

Picking a truck snapshots its loads (stop order, ETA, volume, contact) so the list stays
as it was requested even if the plan is edited later; the planner then sets each
shipment's status (requested / pending / confirmed / declined) on the Appointments page.
Rows are kept as history, so the confirmed and declined share can be tracked over time.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from typing import Any

from routeopt.domain import SITE_BY_CODE, Plan, Stop, clock

STATUSES = ["Not requested", "Requested", "Pending", "Confirmed", "Declined"]
# a stop's outcome: confirmed when every shipment is, declined (the customer can't take
# it that day; dropped from the route) when any is, else still open
OUTCOMES = ["Confirmed", "Declined", "Open"]


def appointment_rows(plan: Plan, stops: list[Stop], truck_number: str) -> list[dict[str, Any]]:
    """One row per shipment on the truck's loads, in load / stop order."""
    by_id = {s.stop_id: s for s in stops}
    truck = plan.truck(truck_number)
    rows: list[dict[str, Any]] = []
    for load in truck.loads:
        for visit in load.stops:
            stop = by_id.get(visit.stop_id)
            if stop is None:
                continue
            for sh in stop.shipments:
                rows.append(
                    {
                        "run_id": plan.run_id,
                        "version": plan.plan_version,
                        "truck": truck.number,
                        "load_seq": load.seq,
                        "stop_seq": visit.seq,
                        "stop_id": stop.stop_id,
                        "shipment_id": sh.shipment_id,
                        "customer": sh.customer_name,
                        "address": f"{stop.address}, {stop.city} {stop.zip}",
                        "bill_to_code": sh.bill_to_code,
                        "delivery_date": plan.plan_date,
                        "eta": clock(visit.arrive_s),
                        "window": f"{clock(visit.eta_window_start_s)} – {clock(visit.eta_window_end_s)}",
                        "volume_cuft": round(sh.volume_cuft, 2),
                        "miles_from_prev": visit.dist_mi_from_prev,
                        "stop_transit_min": visit.transit_min_from_prev,
                        "poc": stop.contact,
                        "dc": load.site,
                    }
                )
    return rows


def stop_outcome(statuses: list[str]) -> str:
    if statuses and all(s == "Confirmed" for s in statuses):
        return "Confirmed"
    if any(s == "Declined" for s in statuses):
        return "Declined"
    return "Open"


def stop_outcomes(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """One row per picked stop: delivery day, truck, volume and its outcome."""
    stops: dict[tuple[Any, ...], dict[str, Any]] = {}
    for r in rows:
        key = (r["run_id"], r["truck"], r["load_seq"], r["stop_seq"])
        st = stops.setdefault(
            key,
            {
                "delivery_date": r["delivery_date"],
                "run_id": r["run_id"],
                "truck": r["truck"],
                "load_seq": r["load_seq"],
                "stop_seq": r["stop_seq"],
                "customer": r["customer"],
                "volume_cuft": 0.0,
                "statuses": [],
            },
        )
        st["volume_cuft"] += r["volume_cuft"] or 0.0
        st["statuses"].append(r["status"])
    out = []
    for st in stops.values():
        st["outcome"] = stop_outcome(st.pop("statuses"))
        out.append(st)
    return out


def outcome_summary(stops: list[dict[str, Any]]) -> dict[str, dict[str, float]]:
    """Per outcome: stops, % of stops, cu ft, % of cu ft."""
    n = len(stops) or 1
    vol = sum(s["volume_cuft"] for s in stops) or 1.0
    out: dict[str, dict[str, float]] = {}
    for o in OUTCOMES:
        sel = [s for s in stops if s["outcome"] == o]
        v = sum(s["volume_cuft"] for s in sel)
        out[o] = {
            "stops": len(sel),
            "stops_pct": round(100 * len(sel) / n, 1),
            "volume": round(v, 1),
            "volume_pct": round(100 * v / vol, 1),
        }
    return out


def stop_status(statuses: list[str]) -> str:
    """One status for a stop: the shared one, else the least advanced (so a stop shows
    Confirmed only when every shipment is; any Declined wins)."""
    if "Declined" in statuses:
        return "Declined"
    if len(set(statuses)) == 1:
        return statuses[0]
    return min(statuses, key=STATUSES.index)


def group_by_stop(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Shipment rows -> one row per stop (same truck, load and stop), keeping the row ids
    so a status set on the stop can be written to each of its shipments."""
    stops: dict[tuple[Any, ...], dict[str, Any]] = {}
    for r in rows:
        key = (r["run_id"], r["version"], r["truck"], r["load_seq"], r["stop_seq"])
        st = stops.get(key)
        if st is None:
            st = stops[key] = {
                **{
                    k: r[k]
                    for k in r
                    if k
                    not in ("id", "shipment_id", "volume_cuft", "load_id", "appt_num", "appt_time")
                },
                "ids": [],
                "shipment_ids": [],
                "volume_cuft": 0.0,
                "statuses": [],
                "notes": [],
                "bill_tos": [],
                "load_ids": [],
                "appts": [],
                "times": [],
            }
        st["ids"].append(r["id"])
        st["shipment_ids"].append(r["shipment_id"])
        st["volume_cuft"] += r["volume_cuft"] or 0.0
        st["statuses"].append(r["status"])
        if r.get("note") and r["note"] not in st["notes"]:
            st["notes"].append(r["note"])
        if r.get("bill_to_code") and r["bill_to_code"] not in st["bill_tos"]:
            st["bill_tos"].append(r["bill_to_code"])
        for key, bucket in (("load_id", "load_ids"), ("appt_num", "appts"), ("appt_time", "times")):
            if r.get(key) and r[key] not in st[bucket]:
                st[bucket].append(r[key])
    out = []
    for st in stops.values():
        st["shipments"] = len(st["ids"])
        st["volume_cuft"] = round(st["volume_cuft"], 1)
        st["status"] = stop_status(st.pop("statuses"))
        st["note"] = " | ".join(st.pop("notes"))
        st["bill_to_code"] = ", ".join(st.pop("bill_tos")) or None
        st["load_id"] = ", ".join(st.pop("load_ids")) or None
        st["appt_num"] = ", ".join(st.pop("appts")) or None
        times = st.pop("times")
        st["appt_time"] = times[0] if times else None
        st["shipment_ids"] = ", ".join(st["shipment_ids"])
        out.append(st)
    return out


WAREHOUSE_CODES = {"LGC2": "N1F", "LGC7": "NF4", "LGC8": "NF5"}  # TMS pick-up codes


def origin_label(dc: str) -> str:
    code = WAREHOUSE_CODES.get(dc)
    return f"{dc} ({code})" if code else dc


def warehouse_label(dc: str) -> str:
    code = WAREHOUSE_CODES.get(dc)
    return f"{dc} WAREHOUSE ({code})" if code else f"{dc} WAREHOUSE"


def warehouse_address(dc: str) -> str:
    site = SITE_BY_CODE.get(dc)
    return site.address if site else ""


def time_text(hhmm: str | None) -> str:
    """'13:05' -> '1:05 PM' (blank stays blank)."""
    if not hhmm:
        return ""
    h, m = (int(x) for x in hhmm.split(":")[:2])
    return f"{(h % 12) or 12}:{m:02d} {'AM' if h < 12 else 'PM'}"


# the export: stop 1 is the warehouse (the load's total volume), then the customer stops
EXPORT_COLUMNS = [
    "load ID",
    "Stop",
    "date",
    "APPT time",
    "APPT #",
    "Customer",
    "Address",
    "Volume",
    "Miles from prev",
    "POC",
    "Note",
]


def load_sheet_rows(stops: list[dict[str, Any]], exclude_declined: bool = True) -> list[dict]:
    """One load's export rows from its ``group_by_stop`` rows: the warehouse as stop 1,
    then the customer stops renumbered 2..n (declined stops left out by default)."""
    kept = [s for s in stops if not (exclude_declined and s["status"] == "Declined")]
    if not stops:
        return []
    first = stops[0]
    load_id = next((s["load_id"] for s in stops if s.get("load_id")), None) or ""
    total = round(sum(s["volume_cuft"] for s in kept), 1)
    out = [
        {
            "load ID": load_id,
            "Stop": 1,
            "date": first["delivery_date"],
            "APPT time": "",
            "APPT #": "",
            "Customer": warehouse_label(first["dc"]),
            "Address": warehouse_address(first["dc"]),
            "Volume": f"Load total volume : {total:,.1f}",
            "Miles from prev": None,
            "POC": "",
            "Note": "",
        }
    ]
    for n, s in enumerate(kept, start=2):
        out.append(
            {
                "load ID": s.get("load_id") or load_id,
                "Stop": n,
                "date": s["delivery_date"],
                "APPT time": time_text(s.get("appt_time")),
                "APPT #": s.get("appt_num") or "",
                "Customer": s["customer"],
                "Address": s["address"],
                "Volume": round(s["volume_cuft"], 1),
                "Miles from prev": s["miles_from_prev"],
                "POC": s["poc"] or "",
                "Note": s["note"] or "",
            }
        )
    return out


def export_stops_by_load(
    rows: list[dict[str, Any]],
    exclude_declined: bool = True,
    only: tuple[str, str, int] | None = None,
) -> bytes:
    """Excel of every load in ``rows`` (or only (run_id, truck, load_seq)): per load the
    warehouse as stop 1, then its stops; a blank row between loads."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Loads"
    ws.append(EXPORT_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center", vertical="center")
    loads: dict[tuple[str, str, int], list[dict[str, Any]]] = {}
    for st in group_by_stop(rows):
        loads.setdefault((st["run_id"], st["truck"], st["load_seq"]), []).append(st)
    keys = [k for k in loads if only is None or k == only]
    for n, key in enumerate(keys):
        if n:
            ws.append([])
        for i, r in enumerate(load_sheet_rows(loads[key], exclude_declined)):
            ws.append([r[c] for c in EXPORT_COLUMNS])
            ws.cell(ws.max_row, EXPORT_COLUMNS.index("date") + 1).number_format = "m/d/yyyy"
            if i == 0:  # the warehouse row
                for c in ws[ws.max_row]:
                    c.font = Font(bold=True)
                    c.fill = PatternFill("solid", fgColor="E7ECF5")
    for i, w in enumerate([12, 6, 11, 11, 14, 32, 42, 24, 9, 26, 28], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def manual_rows(
    load_rows: list[dict[str, Any]],
    stops: list[Stop],
    shipment_ids: str,
    *,
    manual: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Rows that add a stop to the load whose rows are ``load_rows`` (one truck, one load,
    one stage).  ``shipment_ids`` = one or more IDs (comma / space separated); each brings
    its whole stop from the run's upload (every shipment at that address, details copied).
    A stop the load already has is filled up; otherwise it becomes the new last stop.  A
    shipment that is not in the upload needs ``manual`` (customer, address, volume_cuft,
    poc) and becomes its own last stop."""
    if not load_rows:
        raise ValueError("The load has no rows to add to.")
    wanted = [x for x in re.split(r"[,\s]+", shipment_ids.strip()) if x]
    if not wanted:
        raise ValueError("Enter a shipment ID.")
    base = load_rows[0]
    on_load = {r["shipment_id"] for r in load_rows}
    live = [r for r in load_rows if r["stop_seq"] < 900]
    last = max((r["stop_seq"] for r in live), default=0)
    common: dict[str, Any] = {
        k: base[k] for k in ("run_id", "version", "truck", "load_seq", "delivery_date", "dc")
    }
    common |= {
        "eta": "",
        "window": "",
        "miles_from_prev": 0.0,
        "stop_transit_min": 0.0,
        "status": STATUSES[0],
        "note": "added manually",
        "load_id": next((r["load_id"] for r in load_rows if r.get("load_id")), None),
        "stage": base.get("stage"),
    }
    out: list[dict[str, Any]] = []
    seq_of: dict[str, int] = {r["stop_id"]: r["stop_seq"] for r in live}
    done: set[str] = set()
    for sid in wanted:
        if sid in on_load or sid in done:
            continue
        stop = next((s for s in stops if sid in s.shipment_ids), None)
        if stop is None:
            manual = manual or {}
            if not manual.get("customer") or not manual.get("address"):
                raise ValueError(
                    f"{sid} is not in this run's upload: enter its customer and address to add it."
                )
            last += 1
            out.append(
                common
                | {
                    "shipment_id": sid,
                    "stop_id": f"M-{sid}"[:32],
                    "stop_seq": last,
                    "customer": str(manual["customer"]),
                    "address": str(manual["address"]),
                    "bill_to_code": None,
                    "volume_cuft": float(manual.get("volume_cuft") or 0.0),
                    "poc": str(manual.get("poc") or ""),
                }
            )
            done.add(sid)
            continue
        if stop.stop_id not in seq_of:
            last += 1
            seq_of[stop.stop_id] = last
        for sh in stop.shipments:  # the whole stop
            if sh.shipment_id in on_load or sh.shipment_id in done:
                continue
            out.append(
                common
                | {
                    "shipment_id": sh.shipment_id,
                    "stop_id": stop.stop_id,
                    "stop_seq": seq_of[stop.stop_id],
                    "customer": sh.customer_name,
                    "address": f"{stop.address}, {stop.city} {stop.zip}",
                    "bill_to_code": sh.bill_to_code,
                    "volume_cuft": round(sh.volume_cuft, 2),
                    "poc": stop.contact,
                }
            )
            done.add(sh.shipment_id)
    if not out:
        raise ValueError("Everything entered is already on this load.")
    return out


def finalized_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The loads that go: per run the latest set (re-optimized if there is one, else the
    original), keeping only the stops whose every shipment is Confirmed."""
    reopt_runs = {r["run_id"] for r in rows if r.get("stage") == "reoptimized"}
    current = [
        r
        for r in rows
        if (r.get("stage") == "reoptimized") == (r["run_id"] in reopt_runs)
        and r["stop_seq"] < 900
    ]
    statuses: dict[tuple[Any, ...], list[str]] = {}
    for r in current:
        key = (r["run_id"], r["truck"], r["load_seq"], r["stop_seq"])
        statuses.setdefault(key, []).append(r["status"])
    return [
        {**r, "stage": "finalized"}
        for r in current
        if stop_outcome(statuses[(r["run_id"], r["truck"], r["load_seq"], r["stop_seq"])])
        == "Confirmed"
    ]


def period_key(day: dt.date, period: str) -> tuple[dt.date, str]:
    """(sort key, label) of the day / week (Mon-Sun) / month / year a delivery day is in."""
    if period == "Week":
        start = day - dt.timedelta(days=day.weekday())
        return start, f"Wk {start:%m/%d}"
    if period == "Month":
        return day.replace(day=1), f"{day:%Y-%m}"
    if period == "Year":
        return day.replace(month=1, day=1), f"{day:%Y}"
    return day, f"{day:%m/%d}"
