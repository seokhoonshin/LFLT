"""Load summary export (TMS load-list layout) and the floating driver start."""

from __future__ import annotations

import datetime as dt

from openpyxl import load_workbook
from routeopt.export.excel import export_excel
from routeopt.export.load_summary import (
    COLUMNS,
    export_load_summary_pdf,
    load_summary_rows,
    truck_code,
)
from routeopt.pipeline import default_trucks, optimize
from routeopt.travel import node_coords
from routeopt.travel.haversine import HaversineProvider

from tests.test_ca_lflt import _ca, _ship, _stops

FAST = HaversineProvider()


def _plan(tmp_path):  # type: ignore[no-untyped-def]
    s = _ca()
    ships = [_ship(15, 30 * i, address=f"{i} Rialto Rd", volume=60) for i in range(3)]
    ships += [
        _ship(15, 30 * i + 15, site="LGC8", address=f"{i} Fontana Rd", volume=60, pallet=True)
        for i in range(2)
    ]
    ships = [x.model_copy(update={"declared_value": 10.5}) for x in ships]
    stops = _stops(ships, s)
    res = optimize(stops, FAST.matrix(node_coords(stops)), s, default_trucks()[:2])
    return res.plan, stops, s


def test_truck_code() -> None:
    assert truck_code("T01", 5) == "UFPL_STL_TRK_01"
    assert truck_code("Truck", 3) == "UFPL_STL_TRK_03"


def test_load_summary_rows_and_exports(tmp_path) -> None:  # type: ignore[no-untyped-def]
    plan, stops, s = _plan(tmp_path)
    rows = load_summary_rows(plan, stops, s)
    loads = [ld for t in plan.trucks for ld in t.loads if ld.stops]
    assert len(rows) == len(loads)
    for row, load in zip(rows, loads, strict=True):
        assert list(row) == COLUMNS
        assert row["LOAD_ID"] == "" and row["CARR_CD"] == "UFPL" and row["SERVICE CODE"] == "LFLT"
        assert row["PICK_UP"] == {"LGC2": "N1F", "LGC8": "NF5"}[load.site]
        assert row["REMARK"].startswith(
            "STL - LIVELOAD - " + ("PALLET" if load.pallet else "FLOORLOAD")
        )
        assert (" STOP" in row["REMARK"]) == (len(load.stops) > 1)
        assert row["# PIECES"] == load.load_out_pieces
        n_ship = sum(
            len(next(x for x in stops if x.stop_id == v.stop_id).shipments) for v in load.stops
        )
        assert row["DECLARED VAL"] == round(10.5 * n_ship, 2)
        assert row["DROP_APT"].date() == plan.plan_date and row["PICK_APT"] < row["DROP_APT"]
    xlsx = export_excel(plan, stops, s, tmp_path)
    ws = load_workbook(xlsx)["Load Summary"]
    assert [c.value for c in ws[1]] == COLUMNS and ws.max_row == len(rows) + 1
    assert isinstance(ws.cell(row=2, column=COLUMNS.index("PICK_APT") + 1).value, dt.datetime)
    pdf = export_load_summary_pdf(plan, stops, s, tmp_path)
    assert pdf.read_bytes()[:4] == b"%PDF"


def test_results_load_table_starts_with_the_warehouse_pickup(tmp_path) -> None:  # type: ignore[no-untyped-def]
    import sys
    from pathlib import Path

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))
    from routeopt.domain import clock

    from _common import load_table

    plan, stops, _s = _plan(tmp_path)
    truck = next(t for t in plan.trucks if any(ld.stops for ld in t.loads))
    load = next(ld for ld in truck.loads if ld.stops)
    df = load_table(plan, truck.number, load.seq, {s.stop_id: s for s in stops})
    first = df.iloc[0]
    assert first["#"] == "DC" and first["Depart"] == clock(load.depart_s)
    assert "warehouse pickup" in first["Customer"] and len(df) == len(load.stops) + 1
