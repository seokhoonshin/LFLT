"""M6 tests: Excel workbook (incl. round trip), PDF route sheets, map."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from routeopt.edit.revalidate import set_export_override
from routeopt.export.excel import ROUTES_COLUMNS, export_excel
from routeopt.export.mapping import build_map, map_html, truck_colors
from routeopt.export.route_sheets_pdf import export_route_sheets
from routeopt.ingest.consolidate import consolidate
from routeopt.ingest.excel_loader import apply_mapping, auto_map_columns, read_table
from routeopt.ingest.validator import validate_rows
from routeopt.pipeline import OptimizeResult, default_trucks, optimize
from routeopt.travel import node_coords
from routeopt.travel.haversine import HaversineProvider

from tests.factories import (
    BIG_SITE,
    RETURN_SITE,
    make_settings,
    make_shipment,
    make_stops,
    renumber,
)


@pytest.fixture(scope="module")
def solved() -> OptimizeResult:
    settings = make_settings(solve_time_seconds=3)
    big = make_shipment(
        address="1 Big Builder Rd",
        lat=BIG_SITE[0],
        lon=BIG_SITE[1],
        volume=600,
        service=120,
        order_type="BUILDER",
    )
    rtn = make_shipment(
        "RTN", address="9 Return Ln", lat=RETURN_SITE[0], lon=RETURN_SITE[1], volume=40, service=20
    ).model_copy(update={"notes": "Haul away"})
    inc = make_shipment(
        "INC", address="9 Return Ln", lat=RETURN_SITE[0], lon=RETURN_SITE[1], service=30
    )
    small = make_stops(6, volume=30, service=30, max_miles=8, prefix="B")
    stops = renumber(consolidate([big, rtn, inc], settings).stops + small)
    matrix = HaversineProvider().matrix(node_coords(stops))
    res = optimize(stops, matrix, settings, default_trucks()[:2])
    assert res.plan.kpis.loads >= 2
    return res


def test_excel_export_sheets_and_round_trip(solved: OptimizeResult, tmp_path: Path) -> None:
    settings = make_settings(solve_time_seconds=3)
    plan = solved.plan
    path = export_excel(
        plan,
        solved.stops,
        settings,
        tmp_path,
        trucks_snapshot=default_trucks()[:2],
        source_name="s.xlsx",
    )
    assert path.name == f"routeopt_{plan.plan_date.isoformat()}_{plan.run_id}_v0.xlsx"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Routes", "Stops", "Unassigned", "Trucks", "Settings"]
    routes = wb["Routes"]
    assert [c.value for c in routes[1]] == ROUTES_COLUMNS
    served_shipments = [
        sh.shipment_id
        for t in plan.trucks
        for load in t.loads
        for v in load.stops
        for sh in next(s for s in solved.stops if s.stop_id == v.stop_id).shipments
    ]
    assert routes.max_row - 1 == len(served_shipments)
    assert routes.freeze_panes == "A2"
    eta_col = ROUTES_COLUMNS.index("eta") + 1
    assert routes.cell(row=2, column=eta_col).number_format == "h:mm AM/PM"
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "ESTIMATED" in summary_text and "Trucks used" in summary_text
    # acceptance 13: the Routes sheet re-imports cleanly through the loader
    table = read_table(path, sheet_name="Routes")
    mapping = auto_map_columns(table.headers)
    assert mapping.complete, mapping.missing_required
    result = validate_rows(apply_mapping(table, mapping.mapping), plan.plan_date, settings)
    assert result.ok, result.errors[:3]
    assert sorted(s.shipment_id for s in result.shipments) == sorted(served_shipments)
    stops_ws = wb["Stops"]
    assert stops_ws.max_row - 1 == plan.kpis.stops_served
    trucks_ws = wb["Trucks"]
    assert trucks_ws.max_row - 1 == len(plan.trucks)


def test_excel_unassigned_and_override(solved: OptimizeResult, tmp_path: Path) -> None:
    settings = make_settings(solve_time_seconds=3)
    from routeopt.domain import DayTag, UnassignedStop

    plan = set_export_override(solved.plan, "customer insisted")
    victim = plan.trucks[0].loads[0].stops.pop()
    plan.unassigned.append(
        UnassignedStop(
            stop_id=victim.stop_id,
            tier=1,
            day_tag=DayTag.D,
            volume=30,
            reason="Out of time",
            detail="test",
        )
    )
    plan.out_of_scope.append(
        UnassignedStop(
            stop_id="S999",
            tier=4,
            day_tag=DayTag.D1,
            volume=10,
            reason="Not included: D+1 pull-forward is off",
        )
    )
    path = export_excel(plan, solved.stops, settings, tmp_path, extra_warnings=["extra"])
    wb = load_workbook(path)
    un = wb["Unassigned"]
    values = [c.value for row in un.iter_rows() for c in row if c.value]
    assert (
        victim.stop_id in values
        and "Out of time" in values
        and "Not in scope (pull-forward off)" in values
    )
    summary = " ".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value)
    assert (
        "EXPORTED WITH OVERRIDE" in summary
        and "customer insisted" in summary
        and "extra" in summary
    )


def test_pdf_route_sheets(solved: OptimizeResult, tmp_path: Path) -> None:
    settings = make_settings(solve_time_seconds=3)
    plan = set_export_override(solved.plan, "ok")
    path = export_route_sheets(plan, solved.stops, settings, tmp_path)
    data = path.read_bytes()
    assert data[:4] == b"%PDF" and len(data) > 1500
    pages = data.count(b"/Type /Page") - data.count(b"/Type /Pages")
    assert pages >= plan.kpis.loads + 1  # overview + one per load
    path2 = export_route_sheets(
        plan, solved.stops, settings, tmp_path / "no_overview", include_overview=False
    )
    assert path2.exists()


def test_map_contains_markers_and_layers(solved: OptimizeResult) -> None:
    settings = make_settings(solve_time_seconds=3)
    from routeopt.domain import DayTag, UnassignedStop

    plan = solved.plan.model_copy(deep=True)
    victim = plan.trucks[0].loads[0].stops.pop()
    plan.unassigned.append(
        UnassignedStop(
            stop_id=victim.stop_id, tier=1, day_tag=DayTag.D, volume=30, reason="Out of time"
        )
    )
    plan.out_of_scope.append(
        UnassignedStop(
            stop_id=plan.trucks[0].loads[0].stops[0].stop_id,
            tier=4,
            day_tag=DayTag.D1,
            volume=1,
            reason="off",
        )
    )
    calls: list[int] = []

    def geometry(coords):  # type: ignore[no-untyped-def]
        calls.append(len(coords))
        return coords  # straight lines, but recorded

    html = map_html(plan, solved.stops, settings, geometry_fn=geometry)
    assert "Hub" in html and "T01·L1·#1" in html and "Unassigned" in html and "Not in scope" in html
    assert "leaflet-control-layers" in html or "LayerControl" in html or "layer" in html.lower()
    assert calls and all(n >= 3 for n in calls)
    colors = truck_colors(plan)
    assert colors["T01"] != colors["T02"]
    m = build_map(plan, solved.stops, settings)
    assert m is not None
