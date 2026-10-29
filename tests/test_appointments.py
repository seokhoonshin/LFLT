"""Appointments: picking a truck on Results and tracking each shipment's status."""

from __future__ import annotations

from routeopt.appointments import STATUSES, appointment_rows
from routeopt.db.repo import Repository

from tests.test_load_summary import _plan


def test_picked_truck_is_listed_and_statuses_are_saved(tmp_path) -> None:  # type: ignore[no-untyped-def]
    plan, stops, _s = _plan(tmp_path)
    stops = [
        s.model_copy(
            update={
                "shipments": [sh.model_copy(update={"bill_to_code": "BT1"}) for sh in s.shipments]
            }
        )
        for s in stops
    ]
    truck = next(t for t in plan.trucks if any(load.stops for load in t.loads))
    rows = appointment_rows(plan, stops, truck.number)
    n_ship = sum(
        len(next(s for s in stops if s.stop_id == v.stop_id).shipments)
        for load in truck.loads
        for v in load.stops
    )
    assert len(rows) == n_ship
    first = rows[0]
    assert first["truck"] == truck.number and first["load_seq"] == 1 and first["stop_seq"] == 1
    assert first["bill_to_code"] == "BT1" and first["delivery_date"] == plan.plan_date
    assert first["stop_transit_min"] == 0 and first["eta"]

    repo = Repository(f"sqlite:///{tmp_path / 'a.db'}")
    repo.init_schema()
    assert repo.select_truck_for_appointments(rows) == n_ship
    assert repo.select_truck_for_appointments(rows) == 0  # picking again adds nothing
    assert repo.selected_trucks(plan.run_id) == {truck.number}
    listed = repo.list_appointments()
    assert [r["status"] for r in listed] == [STATUSES[0]] * n_ship
    changed = repo.update_appointments({listed[0]["id"]: {"status": "Confirmed", "note": "ok"}})
    assert changed == 2
    again = repo.list_appointments()[0]
    assert again["status"] == "Confirmed" and again["note"] == "ok"
    repo.unselect_truck(plan.run_id, truck.number)
    assert repo.list_appointments() == [] and repo.selected_trucks(plan.run_id) == set()


def test_stop_outcomes_and_summary() -> None:
    import datetime as dt

    from routeopt.appointments import outcome_summary, stop_outcomes

    day = dt.date(2026, 10, 30)

    def row(stop: int, status: str, vol: float) -> dict:
        return {
            "run_id": "r",
            "truck": "T01",
            "load_seq": 1,
            "stop_seq": stop,
            "delivery_date": day,
            "customer": "c",
            "volume_cuft": vol,
            "status": status,
        }

    rows = [
        row(1, "Confirmed", 100),
        row(1, "Confirmed", 50),
        row(2, "Confirmed", 10),
        row(2, "Declined", 30),  # one declined shipment drops the stop
        row(3, "Pending", 60),
        row(4, "Confirmed", 50),
    ]
    stops = {s["stop_seq"]: s for s in stop_outcomes(rows)}
    assert [stops[i]["outcome"] for i in (1, 2, 3, 4)] == [
        "Confirmed",
        "Declined",
        "Open",
        "Confirmed",
    ]
    summary = outcome_summary(list(stops.values()))
    assert summary["Confirmed"] == {
        "stops": 2,
        "stops_pct": 50.0,
        "volume": 200.0,
        "volume_pct": 66.7,
    }
    assert summary["Declined"]["stops"] == 1 and summary["Declined"]["stops_pct"] == 25.0
    assert summary["Open"]["volume"] == 60.0


def test_group_by_stop_merges_shipments_of_one_stop() -> None:
    import datetime as dt

    from routeopt.appointments import group_by_stop

    base = {
        "run_id": "r",
        "version": 0,
        "truck": "T01",
        "load_seq": 1,
        "delivery_date": dt.date(2026, 11, 2),
        "customer": "C",
        "address": "1 A St",
        "eta": "8:00 AM",
        "window": "",
        "miles_from_prev": 5.0,
        "stop_transit_min": 0.0,
        "poc": "",
        "dc": "LGC2",
        "note": "",
    }
    rows = [
        {**base, "id": 1, "stop_seq": 1, "shipment_id": "A", "volume_cuft": 10.0,
         "bill_to_code": "B1", "status": "Confirmed"},
        {**base, "id": 2, "stop_seq": 1, "shipment_id": "B", "volume_cuft": 5.5,
         "bill_to_code": "B1", "status": "Requested"},
        {**base, "id": 3, "stop_seq": 2, "shipment_id": "C", "volume_cuft": 7.0,
         "bill_to_code": None, "status": "Declined"},
    ]  # fmt: skip
    stops = group_by_stop(rows)
    assert len(stops) == 2
    first, second = stops
    assert first["ids"] == [1, 2] and first["shipments"] == 2
    assert first["shipment_ids"] == "A, B" and first["volume_cuft"] == 15.5
    assert first["bill_to_code"] == "B1"
    assert first["status"] == "Requested"  # not every shipment is confirmed yet
    assert second["status"] == "Declined" and second["bill_to_code"] is None


def _rows_for_export() -> list[dict]:
    import datetime as dt

    base = {
        "run_id": "r",
        "version": 0,
        "truck": "T01",
        "load_seq": 1,
        "delivery_date": dt.date(2026, 11, 2),
        "address": "1 A St",
        "eta": "8:00 AM",
        "window": "",
        "miles_from_prev": 5.0,
        "stop_transit_min": 0.0,
        "poc": "",
        "dc": "LGC2",
        "note": "",
        "bill_to_code": None,
        "load_id": "19007709",
        "appt_num": None,
        "appt_time": None,
    }
    return [
        {**base, "id": 1, "stop_seq": 1, "customer": "A", "shipment_id": "a1",
         "volume_cuft": 10.0, "status": "Confirmed", "appt_time": "08:00"},
        {**base, "id": 2, "stop_seq": 1, "customer": "A", "shipment_id": "a2",
         "volume_cuft": 5.0, "status": "Confirmed", "appt_time": "08:00"},
        {**base, "id": 3, "stop_seq": 2, "customer": "B", "shipment_id": "b1",
         "volume_cuft": 7.0, "status": "Declined"},
        {**base, "id": 4, "stop_seq": 3, "customer": "C", "shipment_id": "c1",
         "volume_cuft": 3.0, "status": "Requested", "appt_num": "APT-9"},
    ]  # fmt: skip


def test_export_stops_by_load_leaves_out_declined() -> None:
    import io

    from openpyxl import load_workbook
    from routeopt.appointments import EXPORT_COLUMNS, export_stops_by_load

    ws = load_workbook(io.BytesIO(export_stops_by_load(_rows_for_export()))).active
    header = [c.value for c in ws[1]]
    assert header == EXPORT_COLUMNS
    assert header == [
        "load ID",
        "Stop",
        "date",
        "APPT time",
        "APPT #",
        "Customer",
        "Address",
        "Volume",
        "Miles from prev",
        "POC",
        "Note",
    ]
    body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    col = {h: i for i, h in enumerate(header)}
    wh = body[0]  # stop 1 is the warehouse with the load's total volume
    assert wh[col["Stop"]] == 1 and wh[col["Customer"]] == "LGC2 WAREHOUSE (N1F)"
    assert wh[col["Address"]] == "5550 Mango Ave, Rialto, CA 92377"
    assert wh[col["Volume"]] == "Load total volume : 18.0"  # declined B not counted
    assert wh[col["APPT #"]] is None and wh[col["load ID"]] == "19007709"
    stops = body[1:]
    assert [r[col["Customer"]] for r in stops] == ["A", "C"]  # B declined: left out
    assert [r[col["Stop"]] for r in stops] == [2, 3]
    assert stops[0][col["Volume"]] == 15.0 and stops[0][col["APPT time"]] == "8:00 AM"
    assert stops[1][col["APPT #"]] == "APT-9"
    keep = export_stops_by_load(_rows_for_export(), exclude_declined=False)
    ws2 = load_workbook(io.BytesIO(keep)).active
    assert sum(r[col["Customer"]].value == "B" for r in ws2.iter_rows(min_row=2)) == 1
    one = export_stops_by_load(_rows_for_export(), only=("r", "T01", 1))
    assert load_workbook(io.BytesIO(one)).active.max_row == 4
    none = export_stops_by_load(_rows_for_export(), only=("r", "T09", 1))
    assert load_workbook(io.BytesIO(none)).active.max_row == 1


def test_old_database_gets_new_appointment_columns(tmp_path) -> None:  # type: ignore[no-untyped-def]
    import sqlite3

    db = tmp_path / "old.db"
    con = sqlite3.connect(db)
    con.execute("create table appointments (id integer primary key, run_id varchar(64))")
    con.commit()
    con.close()
    repo = Repository(f"sqlite:///{db}")
    repo.init_schema()
    cols = {r[1] for r in sqlite3.connect(db).execute("pragma table_info(appointments)")}
    assert {"load_id", "appt_num", "appt_time"} <= cols


def test_stop_order_and_appointment_time_are_saved(tmp_path) -> None:  # type: ignore[no-untyped-def]
    plan, stops, _s = _plan(tmp_path)
    truck = next(t for t in plan.trucks if any(load.stops for load in t.loads))
    repo = Repository(f"sqlite:///{tmp_path / 'b.db'}")
    repo.init_schema()
    repo.select_truck_for_appointments(appointment_rows(plan, stops, truck.number))
    rows = repo.list_appointments()
    first = [r for r in rows if r["stop_seq"] == 1 and r["load_seq"] == rows[0]["load_seq"]]
    repo.update_appointments(
        {r["id"]: {"stop_seq": 9, "appt_time": "10:30", "load_id": "123"} for r in first}
    )
    again = [r for r in repo.list_appointments() if r["id"] in {x["id"] for x in first}]
    assert all(r["stop_seq"] == 9 and r["appt_time"] == "10:30" for r in again)
    assert all(r["load_id"] == "123" for r in again)


def test_map_of_one_load_draws_only_that_load(tmp_path) -> None:  # type: ignore[no-untyped-def]
    from routeopt.export.mapping import map_html

    plan, stops, s = _plan(tmp_path)
    loads = [(t.number, ld.seq) for t in plan.trucks for ld in t.loads if ld.stops]
    number, seq = loads[0]
    one = map_html(plan, stops, s, only={(number, seq)})
    assert f"{number}·L{seq}·#1" in one
    others = [x for x in loads if x[0] != number]
    if others:
        assert f"{others[0][0]}·L{others[0][1]}·#1" not in one
